#!/usr/bin/env python3
"""Build the database spreadsheet (one sheet per table) from schema/tables.json.

Usage:
    python3 tools/build_database_sheet.py \
        --equipment local-data/ACC_Downtime_Equipment_Master_v2.xlsx \
        --reasons   local-data/ACC_Downtime_Reason_Master.xlsx \
        --out       local-data/ACC_Downtime_Database.xlsx

The master inputs are real plant data and live in the git-ignored local-data/
folder. Without them the sheet is still built, with empty master tables.
"""
import argparse
import json
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
SCHEMA = json.loads((ROOT / "schema" / "tables.json").read_text(encoding="utf-8"))
TABLES = {t["name"]: t for t in SCHEMA["tables"]}

HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1E5BB8")
TEXT_TYPES = {"text", "date", "timestamp", "enum"}


def clean(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v if v else None
    return v


def read_sheet(path, sheet):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    head = [c.value for c in ws[1]]
    return [dict(zip(head, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]


def build_rows(args):
    rows = {t: [] for t in TABLES}
    rows["lines"] = [dict(r) for r in SCHEMA["lineSeeds"]]
    rows["areas"] = [dict(r, active=True) for r in SCHEMA["areaSeeds"]]
    rows["classifications"] = [dict(r) for r in SCHEMA["classificationSeeds"]]
    rows["roles"] = [dict(r) for r in SCHEMA["roleSeeds"]]
    for role, tabs in SCHEMA["tabAccessDefaults"].items():
        for tab, level in tabs.items():
            rows["tab_access"].append(dict(subject_type="role", subject_id=role, tab_id=tab, level=level))
    rows["settings"] = [dict(r) for r in SCHEMA["settingsDefaults"]]

    area_ids = {a["area_id"] for a in rows["areas"]}

    if args.equipment and Path(args.equipment).exists():
        for r in read_sheet(args.equipment, "EQUIPMENT_MASTER"):
            desc = clean(r.get("Equipment_Description"))
            crit = clean(r.get("Criticality"))
            areas = [t for t in re.split(r"\s*[,/]\s*", str(r.get("Area (downtime group)") or "")) if t]
            review = (not desc) or (not crit) or (not areas) or "REVIEW" in areas
            rows["equipment"].append(dict(
                equipment_id=str(r["Equipment_ID"]).strip().upper(),
                description=desc,
                equipment_type=clean(r.get("Equipment_Type")),
                criticality=crit,
                parent_id=clean(r.get("Parent_Equipment_ID")),
                status="Active",
                origin=clean(r.get("Origin")),
                platform_plant_area=clean(r.get("Platform_Plant_Area")),
                platform_main_area=clean(r.get("Platform_Main_Area")),
                platform_sub_area=None if r.get("Platform_Sub_Area") is None else str(r.get("Platform_Sub_Area")),
                needs_review=review,
                review_note=clean(r.get("Review_Note")),
            ))
        for r in read_sheet(args.equipment, "EQUIPMENT_AREAS"):
            if r["Area (downtime group)"] in area_ids:
                rows["equipment_areas"].append(dict(
                    equipment_id=str(r["Equipment_ID"]).strip().upper(), area_id=r["Area (downtime group)"]))
        for r in read_sheet(args.equipment, "ALIASES"):
            rows["equipment_aliases"].append(dict(
                alias=re.sub(r"\s+", "", str(r["Code as written"]).upper()),
                equipment_id=str(r["Canonical code"]).strip().upper(), reason=clean(r.get("Reason"))))

    if args.reasons and Path(args.reasons).exists():
        for r in read_sheet(args.reasons, "REASON_MASTER"):
            cls = r["Classification"]
            order = int(r["Order in list"])
            rows["reasons"].append(dict(
                reason_id=f"R-{cls}-{order:03d}", classification=cls,
                text_as_listed=str(r["Reason (as listed)"]), match_key=str(r["Match_Key"]),
                order_no=order, active=True))
    if getattr(args, "events_json", None) and Path(args.events_json).exists():
        rows["events"] = json.loads(Path(args.events_json).read_text(encoding="utf-8"))
    if getattr(args, "batches_json", None) and Path(args.batches_json).exists():
        rows["upload_batches"] = json.loads(Path(args.batches_json).read_text(encoding="utf-8"))
    return rows


def check_integrity(rows):
    problems = []
    eq = {r["equipment_id"] for r in rows["equipment"]}
    if len(eq) != len(rows["equipment"]):
        problems.append("duplicate equipment_id")
    areas = {r["area_id"] for r in rows["areas"]}
    for r in rows["equipment_areas"]:
        if r["equipment_id"] not in eq:
            problems.append(f"equipment_areas: unknown equipment {r['equipment_id']}")
        if r["area_id"] not in areas:
            problems.append(f"equipment_areas: unknown area {r['area_id']}")
    for r in rows["equipment_aliases"]:
        if r["equipment_id"] not in eq:
            problems.append(f"alias {r['alias']} -> unknown {r['equipment_id']}")
    cls = {r["code"] for r in rows["classifications"]}
    for r in rows["reasons"]:
        if r["classification"] not in cls:
            problems.append(f"reason {r['reason_id']}: unknown classification")
    ids = [r["reason_id"] for r in rows["reasons"]]
    if len(ids) != len(set(ids)):
        problems.append("duplicate reason_id")
    return problems


def add_table_sheet(wb, table, data):
    ws = wb.create_sheet(table["name"])
    cols = table["columns"]
    ws.append([c["name"] for c in cols])
    for cell, c in zip(ws[1], cols):
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center")
        note = c.get("description") or ""
        meta = f"{c['type']}{' (required)' if c.get('required') else ''}"
        cell.comment = Comment((note + "\n" if note else "") + meta, "schema")
    for r in data:
        ws.append([r.get(c["name"]) for c in cols])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(ws.max_row, 1)}"
    last = 3000 if table["name"] in ("events", "event_edits", "upload_warnings", "audit_log") else max(ws.max_row + 500, 600)
    for i, c in enumerate(cols, 1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = max(12, min(44, len(c["name"]) + 4))
        if c["type"] in TEXT_TYPES:
            for row in range(2, last + 1):
                ws.cell(row=row, column=i).number_format = "@"
        if c["type"] == "enum":
            dv = DataValidation(type="list", formula1='"' + ",".join(c["enum"]) + '"', allow_blank=not c.get("required"))
            dv.add(f"{letter}2:{letter}{last}")
            ws.add_data_validation(dv)
        if c["type"] == "bool":
            dv = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
            dv.add(f"{letter}2:{letter}{last}")
            ws.add_data_validation(dv)
    return ws


def build(args):
    rows = build_rows(args)
    problems = check_integrity(rows)
    if problems:
        print("Integrity problems:", *problems[:20], sep="\n  ")
        sys.exit(1)

    wb = openpyxl.Workbook()
    readme = wb.active
    readme.title = "README"
    for table in SCHEMA["tables"]:
        add_table_sheet(wb, table, rows[table["name"]])

    d = wb.create_sheet("DICTIONARY")
    d.append(["table", "column", "type", "required", "values", "description"])
    for cell in d[1]:
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
    for t in SCHEMA["tables"]:
        for c in t["columns"]:
            d.append([t["name"], c["name"], c["type"], "yes" if c.get("required") else "", ", ".join(c.get("enum", [])), c.get("description", "")])
    d.freeze_panes = "A2"
    for letter, w in zip("ABCDEF", (20, 22, 11, 9, 40, 80)):
        d.column_dimensions[letter].width = w

    lines = [
        ["ACC Downtime Analysis - database spreadsheet", ""],
        ["Schema version", str(SCHEMA["version"])],
        ["", ""],
        ["Purpose", "Storage for the ACC Downtime Analysis app (test phase). One sheet per table; row 1 of each sheet is the header row."],
        ["How to use", "Upload this file to Google Drive, open it with Google Sheets and choose File > Save as Google Sheets. Then connect it to the Apps Script backend as described in backend/apps-script/README.md."],
        ["Do not", "Rename sheets or header cells, insert columns in the middle, or type into events by hand. Use the app to upload, edit or void events."],
        ["Safe to edit by hand", "Only the master sheets (equipment, equipment_areas, equipment_aliases, reasons, areas, settings) - the app also edits them, with an audit trail."],
        ["Text columns", "Date, timestamp and text columns are formatted as Plain Text so Google Sheets never converts them. Keep them plain text."],
        ["users", "Empty on purpose. The first admin is created by the setup function in the Apps Script project (a one-time password is shown once)."],
        ["notifications", "notifications, notification_state and notification_prefs are empty on purpose: the server fills them (alerts, updates, per-person state and preferences). The notif_* settings and app_url are seeded with their defaults."],
        ["events", (f"{len(rows['events'])} events already loaded (see upload_batches). Add new months by uploading the monthly workbook in the app; rows already here are skipped, never duplicated." if rows["events"] else "Empty on purpose. Fill it by uploading the monthly downtime workbook in the app.")],
        ["", ""],
        ["Seeded rows", ""],
    ]
    for name in ("lines", "areas", "classifications", "reasons", "equipment", "equipment_areas", "equipment_aliases", "roles", "tab_access", "settings"):
        lines.append([name, len(rows[name])])
    lines += [["", ""], ["Rows flagged needs_review", sum(1 for r in rows["equipment"] if r["needs_review"])],
              ["Codes with several areas", len({r["equipment_id"] for r in rows["equipment_areas"] if sum(1 for x in rows["equipment_areas"] if x["equipment_id"] == r["equipment_id"]) > 1})]]
    for row in lines:
        readme.append(row)
    readme["A1"].font = Font(bold=True, size=14)
    readme.column_dimensions["A"].width = 30
    readme.column_dimensions["B"].width = 120
    for row in readme.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print("saved", out)
    for name in TABLES:
        print(f"  {name}: {len(rows[name])} rows")


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--equipment")
    p.add_argument("--reasons")
    p.add_argument("--events-json", help="optional: events already stored (exported from the backend)")
    p.add_argument("--batches-json", help="optional: upload batches that go with the events")
    p.add_argument("--out", required=True)
    build(p.parse_args())
